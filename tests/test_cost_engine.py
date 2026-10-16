import copy
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
import cost_engine as engine

ROOT = Path(__file__).resolve().parents[1]


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.config = json.loads((ROOT / "example_business_config.json").read_text())
        self.receipt = json.loads((ROOT / "example_receipt.json").read_text())
        self.rows, _ = engine.receipt_rows(self.receipt, [], self.config)

    def costs(self, rows=None):
        return engine.compute_costs(engine.latest_prices(self.rows if rows is None else rows), self.config)[0]

    def test_converts_ounces_to_pounds(self):
        cost = self.costs()
        self.assertEqual(cost["status"], "complete")
        self.assertEqual(cost["ingredient_cost_per_unit"], 1.3625)
        self.assertEqual(cost["total_cost_per_unit"], 4.1125)
        self.assertAlmostEqual(engine.conversion_factor("lb", "oz"), 16)
        self.assertAlmostEqual(engine.conversion_factor("l", "ml"), 1000)
        self.assertEqual(engine.conversion_factor("ea", "dozen"), 1/12)
        self.assertEqual(engine.conversion_factor("lbs", "lb"), 1)

    def test_incompatible_package_units_block_margins(self):
        self.rows[0][6] = "cs"
        cost = self.costs()
        self.assertEqual(cost["status"], "incomplete")
        self.assertEqual(cost["unit_errors"][0]["ingredient"], "Pork")
        self.assertIsNone(cost["total_cost_per_unit"])
        self.assertIsNone(cost["tiers"]["dine_in"]["margin_pct"])

    def test_missing_ingredient_blocks_margins(self):
        cost = self.costs(self.rows[1:])
        self.assertEqual(cost["missing_prices"], ["Pork"])
        self.assertIsNone(cost["ingredient_cost_per_unit"])
        for field in ("profit_low", "profit_high", "margin_low_pct", "margin_high_pct"):
            self.assertIsNone(cost["tiers"]["catering"][field])

    def test_duplicate_without_id_survives_reordering_and_alias_changes(self):
        duplicate = copy.deepcopy(self.receipt)
        duplicate["items"].reverse()
        duplicate["items"][0]["quantity"] = 40.0
        self.config["items"]["aliases"]["Pork"]["category"] = "Meat"
        with self.assertRaisesRegex(ValueError, "Duplicate receipt"):
            engine.receipt_rows(duplicate, self.rows, self.config)

    def test_explicit_ids_and_legacy_time_based_ids(self):
        first = copy.deepcopy(self.receipt)
        first["receipt_id"] = "123"
        rows, _ = engine.receipt_rows(first, [], self.config)
        with self.assertRaisesRegex(ValueError, "Duplicate receipt"):
            engine.receipt_rows(first, rows, self.config)
        for row in rows:
            row[10] = "legacy_generated_time_id"
        with self.assertRaisesRegex(ValueError, "Duplicate receipt"):
            engine.receipt_rows(self.receipt, rows, self.config)
        other = copy.deepcopy(first)
        other["merchant"] = "Another supplier"
        self.assertTrue(engine.receipt_rows(other, rows, self.config)[0])

    def test_same_day_last_entry_wins_but_older_dates_do_not(self):
        newer = list(self.rows[0])
        newer[7] = 4
        older = list(newer)
        older[0], older[7] = "2026-10-15", 5
        prices = engine.latest_prices(self.rows + [newer, older])
        self.assertEqual(prices["Pork"]["unit_price"], 4)

    def test_bad_quantities_and_dates_rejected(self):
        for value in (0, -1, "10", float("nan"), float("inf"), None):
            with self.subTest(value=value):
                receipt = copy.deepcopy(self.receipt)
                receipt["items"][0]["quantity"] = value
                with self.assertRaises(ValueError):
                    engine.receipt_rows(receipt, [], self.config)
        self.receipt["date"] = None
        with self.assertRaisesRegex(ValueError, "date"):
            engine.receipt_rows(self.receipt, [], self.config)

    def test_invalid_production_and_batch_block_costs(self):
        self.config["overhead"]["monthly_production"] = 0
        self.config["products"]["recipes"]["Pork Dumplings"]["batch_size"] = 0
        cost = self.costs()
        self.assertEqual(len(cost["validation_errors"]), 2)
        self.assertIsNone(cost["total_cost_per_unit"])

    def test_workbook_roundtrip_duplicate_and_high_pricing(self):
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory) / "purchases.xlsx")
            result = engine.run(str(ROOT / "example_business_config.json"),
                                [str(ROOT / "example_receipt.json")], None, output)
            self.assertEqual(result["ledger_rows"], 5)
            wb = load_workbook(output, data_only=True)
            sheet = wb["Margins"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(sheet.cell(2, headers.index("Catering High Price") + 1).value, 11)
            wb.close()
            repeat = engine.run(str(ROOT / "example_business_config.json"),
                                [str(ROOT / "example_receipt.json")], output, output)
            self.assertEqual(repeat["ledger_rows"], 5)
            self.assertEqual(len(repeat["errors"]), 1)

    def test_incomplete_workbook_shows_issues_and_no_margin(self):
        with tempfile.TemporaryDirectory() as directory:
            path = str(Path(directory) / "incomplete.xlsx")
            prices = engine.latest_prices(self.rows[1:])
            engine.write_workbook(path, self.rows[1:], [], prices,
                                  engine.compute_costs(prices, self.config), self.config)
            wb = load_workbook(path, data_only=True)
            sheet = wb["Margins"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIsNone(sheet.cell(2, headers.index("Dine-In Margin %") + 1).value)
            self.assertEqual(sheet.cell(2, headers.index("Status") + 1).value, "incomplete")
            self.assertIn("Pork", sheet.cell(2, headers.index("Issues") + 1).value)
            wb.close()


if __name__ == "__main__":
    unittest.main()
