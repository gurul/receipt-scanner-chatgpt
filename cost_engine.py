"""
cost_engine.py — Standalone receipt cost engine for ChatGPT's Code Interpreter.

ChatGPT reads the receipt photo and writes the line items to a receipt JSON
file; this script does the deterministic part: maps items to canonical names
using the business config aliases, appends them to the purchases ledger,
and recomputes latest prices, cost per unit, and margins.

Self-contained (stdlib + openpyxl) so it can be uploaded as a GPT knowledge file.
Requires Python 3.9+; workbook I/O requires openpyxl.

Usage:
    python cost_engine.py --config business_config.json \
        --receipt receipt.json [--receipt receipt2.json ...] \
        [--ledger purchases.xlsx|purchases.csv] --out purchases.xlsx

    python cost_engine.py --config business_config.json --ledger purchases.xlsx \
        --out purchases.xlsx          # just recompute costs

    python cost_engine.py --config business_config.json --map "GRD PORK 10LB" "SALT"

Receipt JSON format (same as the main pipeline's extraction output):
    {"merchant": "...", "date": "YYYY-MM-DD", "receipt_id": "..." | null,
     "items": [{"raw_description": "...", "name": "...", "quantity": 1.0,
                "unit": "lb", "unit_price": 2.0, "total_price": 2.0}],
     "subtotal": 0.0, "tax": 0.0, "total": 0.0}

Prints a JSON summary to stdout.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
from datetime import datetime

UNMAPPED_PREFIX = "UNMAPPED: "

# Factors use grams, litres, and individual items as base units. Package sizes
# are deliberately not guessed: a case/bag needs an explicit conversion.
_UNITS = {
    "g": ("mass", 1), "kg": ("mass", 1000),
    "oz": ("mass", 28.349523125), "lb": ("mass", 453.59237),
    "ml": ("volume", .001), "l": ("volume", 1),
    "fl oz": ("volume", .0295735295625), "gal": ("volume", 3.785411784),
    "ea": ("count", 1), "dozen": ("count", 12),
}
_UNIT_ALIASES = {"lbs": "lb", "pound": "lb", "pounds": "lb",
                 "ounces": "oz", "ounce": "oz", "each": "ea", "pc": "ea",
                 "pcs": "ea", "liter": "l", "litre": "l", "gallon": "gal"}


def normalize_unit(unit):
    value = str(unit or "").strip().lower()
    return _UNIT_ALIASES.get(value, value)


def conversion_factor(recipe_unit, purchase_unit):
    """Purchase units needed for one recipe unit; reject unknown conversions."""
    target, source = normalize_unit(recipe_unit), normalize_unit(purchase_unit)
    if target and target == source:
        return 1
    if target in _UNITS and source in _UNITS and _UNITS[target][0] == _UNITS[source][0]:
        return _UNITS[target][1] / _UNITS[source][1]
    raise ValueError(f"Cannot convert {purchase_unit!r} to {recipe_unit!r}")


def receipt_fingerprint(rows):
    """Stable across re-imports, item ordering, numeric formatting and alias edits."""
    items = sorted(json.dumps([str(r[2]).strip().upper(), float(r[5]),
                               normalize_unit(r[6]), float(r[8])]) for r in rows)
    payload = [str(rows[0][0]), str(rows[0][1]).strip().upper(), items]
    return "AUTO-" + hashlib.sha256(json.dumps(payload).encode()).hexdigest()

PURCHASES_COLUMNS = [
    "Date", "Supplier", "Raw Description", "Canonical Name", "Category",
    "Quantity", "Unit", "Unit Price", "Total Price", "Source", "Receipt ID"
]

_DEFAULT_TAB_NAMES = {
    "purchases": "Purchases",
    "latest_prices": "Latest Prices",
    "recipes": "Recipes",
    "margins": "Margins",
    "unmapped": "Unmapped Items",
}


def tab_names(config: dict) -> dict:
    return {**_DEFAULT_TAB_NAMES, **config.get("sheets", {}).get("tab_names", {})}


# ---------------------------------------------------------------------------
# Item mapping
# ---------------------------------------------------------------------------
def map_item(raw_description: str, config: dict) -> dict:
    """Map receipt text to a canonical item via word-boundary alias matching; longest pattern wins."""
    alias_map = config.get("items", {}).get("aliases", {})
    raw_upper = raw_description.upper().strip()

    best_match = None
    best_match_len = 0
    for canonical_name, info in alias_map.items():
        for pattern in info["patterns"]:
            pattern_upper = pattern.upper()
            if re.search(r'\b' + re.escape(pattern_upper) + r'\b', raw_upper):
                if len(pattern_upper) > best_match_len:
                    best_match = {
                        "canonical_name": canonical_name,
                        "category": info["category"],
                        "default_unit": info.get("default_unit", "ea"),
                        "matched_pattern": pattern,
                    }
                    best_match_len = len(pattern_upper)

    if best_match:
        return best_match
    return {
        "canonical_name": f"{UNMAPPED_PREFIX}{raw_description}",
        "category": "Uncategorized",
        "default_unit": "ea",
        "matched_pattern": None,
    }


# ---------------------------------------------------------------------------
# Ledger I/O
# ---------------------------------------------------------------------------
def load_ledger(path: str | None, config: dict) -> list[list]:
    """Load Purchases rows (without header) from .xlsx or .csv. Missing path → empty ledger."""
    if not path:
        return []
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        name = tab_names(config)["purchases"]
        ws = wb[name] if name in wb.sheetnames else wb.worksheets[0]
        rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()

    rows = [r for r in rows if any(str(v).strip() for v in r)]
    if rows and str(rows[0][0]).strip().lower() == "date":
        rows = rows[1:]
    return [_normalize_row(r) for r in rows]


def _normalize_row(row: list) -> list:
    row = list(row) + [""] * (len(PURCHASES_COLUMNS) - len(row))
    row = row[:len(PURCHASES_COLUMNS)]
    if isinstance(row[0], datetime):
        row[0] = row[0].strftime("%Y-%m-%d")
    return row


def receipt_rows(receipt: dict, ledger: list[list], config: dict,
                 source: str = "chatgpt") -> tuple[list[list], dict]:
    """Map a receipt's items and build Purchases rows. Raises ValueError on duplicates/bad input."""
    items = receipt.get("items") or []
    if not items:
        raise ValueError("Receipt must have at least one item")
    for i, item in enumerate(items):
        for field in ("quantity", "total_price"):
            if item.get(field) is None:
                raise ValueError(f"Item {i} missing required field: {field}")

    merchant = receipt.get("merchant") or "Unknown"
    date = receipt.get("date")
    try:
        if not isinstance(date, str) or datetime.strptime(date, "%Y-%m-%d").strftime("%Y-%m-%d") != date:
            raise ValueError()
    except (ValueError, TypeError):
        raise ValueError("Receipt date must be a readable YYYY-MM-DD date")
    receipt_id = receipt.get("receipt_id")

    if receipt_id and any(str(r[10]) == str(receipt_id) and
                          str(r[1]).strip().upper() == merchant.strip().upper()
                          for r in ledger):
        raise ValueError(f"Duplicate receipt: '{receipt_id}' is already in the ledger")

    rows, mapped_items = [], []
    for item in items:
        raw = item.get("raw_description") or item.get("name", "")
        m = map_item(raw, config)
        is_unmapped = m["matched_pattern"] is None
        canonical = m["canonical_name"].removeprefix(UNMAPPED_PREFIX)
        qty = item["quantity"]
        total = item["total_price"]
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or
               not math.isfinite(v) for v in (qty, total)) or qty <= 0:
            raise ValueError("Quantity must be positive and prices must be finite numbers")
        unit_price = item.get("unit_price")
        if unit_price is None:
            unit_price = round(total / qty, 4) if qty else total
        if isinstance(unit_price, bool) or not isinstance(unit_price, (int, float)) or not math.isfinite(unit_price):
            raise ValueError("Unit price must be a finite number")
        rows.append([
            date, merchant, raw, canonical, m["category"], qty,
            item.get("unit") or m["default_unit"], unit_price, total, source, receipt_id,
        ])
        mapped_items.append({"raw": raw, "canonical": canonical, "mapped": not is_unmapped,
                             "category": m["category"], "qty": qty, "unit_price": unit_price,
                             "total": total})

    fingerprint = receipt_fingerprint(rows)
    existing_receipts = {}
    for row in ledger:
        existing_receipts.setdefault((str(row[0]), str(row[1]), str(row[10])), []).append(row)
    if any(receipt_fingerprint(group) == fingerprint for group in existing_receipts.values()):
        raise ValueError("Duplicate receipt: matching merchant, date and line items already recorded")
    receipt_id = receipt_id or fingerprint
    for row in rows:
        row[10] = receipt_id

    summary = {
        "receipt_id": receipt_id, "merchant": merchant, "date": date,
        "items": mapped_items,
        "items_mapped": sum(1 for i in mapped_items if i["mapped"]),
        "items_unmapped": sum(1 for i in mapped_items if not i["mapped"]),
        "total": receipt.get("total"),
    }
    return rows, summary


# ---------------------------------------------------------------------------
# Costing
# ---------------------------------------------------------------------------
def latest_prices(ledger: list[list]) -> dict[str, dict]:
    latest = {}
    for row in ledger:
        date, supplier, _raw, canonical, category, qty, unit, unit_price, total = row[:9]
        try:
            unit_price_num = float(unit_price)
            qty_num = float(qty)
        except (ValueError, TypeError):
            continue
        if not math.isfinite(unit_price_num) or not math.isfinite(qty_num) or unit_price_num < 0 or qty_num <= 0 or category == "Uncategorized":
            continue
        date = str(date)
        if canonical not in latest or date >= latest[canonical]["date"]:
            latest[canonical] = {
                "unit_price": unit_price_num,
                "unit": unit,
                "quantity_purchased": qty_num,
                "total_price": float(total) if total not in ("", None) else unit_price_num * qty_num,
                "date": date,
                "supplier": supplier,
            }
    return latest


def compute_costs(prices: dict, config: dict) -> list[dict]:
    recipes = config.get("products", {}).get("recipes", {})
    overhead_config = config.get("overhead", {})
    pricing_tiers = config.get("pricing", {}).get("tiers", {})
    unit_name = config.get("products", {}).get("unit_name", "unit")

    production = overhead_config.get("monthly_production", 1)
    overhead_per_unit = {}
    total_overhead_per_unit = 0
    for cat in overhead_config.get("cost_categories", []):
        per_unit = cat["monthly_amount"] / production if production > 0 else 0
        overhead_per_unit[cat["name"]] = round(per_unit, 4)
        total_overhead_per_unit += per_unit

    results = []
    for product_name, recipe in recipes.items():
        batch_size = recipe.get("batch_size", recipe.get("batch_rolls", 1))
        batch_cost = 0
        ingredient_costs, missing, unit_errors = [], [], []
        for ingredient, usage in recipe.get("ingredients", {}).items():
            info = prices.get(ingredient)
            if info is None:
                missing.append(ingredient)
                continue
            try:
                factor = conversion_factor(usage["unit"], info["unit"])
            except ValueError as exc:
                unit_errors.append({"ingredient": ingredient, "error": str(exc)})
                continue
            cost = info["unit_price"] * factor * usage["qty"]
            batch_cost += cost
            ingredient_costs.append({
                "ingredient": ingredient, "qty_needed": usage["qty"], "unit": usage["unit"],
                "unit_price": info["unit_price"] * factor, "line_cost": round(cost, 2),
                "price_date": info["date"], "supplier": info["supplier"],
            })

        ingredient_per_unit = batch_cost / batch_size if batch_size > 0 else 0
        total_per_unit = ingredient_per_unit + total_overhead_per_unit
        complete = not missing and not unit_errors and batch_size > 0 and production > 0
        entry = {
            "product": product_name,
            "size": recipe.get("size", ""),
            "batch_size": batch_size,
            "unit_name": unit_name,
            "ingredient_cost_per_unit": round(ingredient_per_unit, 4) if complete else None,
            "overhead_breakdown": dict(overhead_per_unit),
            "total_overhead_per_unit": round(total_overhead_per_unit, 4),
            "total_cost_per_unit": round(total_per_unit, 4) if complete else None,
            "status": "complete" if complete else "incomplete",
            "unit_errors": unit_errors,
            "validation_errors": (["Batch size must be positive"] if batch_size <= 0 else []) +
                                 (["Monthly production must be positive"] if production <= 0 else []),
            "ingredients": ingredient_costs,
            "missing_prices": missing,
            "tiers": {},
        }

        for tier_key, tier in pricing_tiers.items():
            tier_prices = tier.get("prices", {}).get(product_name, {})
            t = {"label": tier.get("label", tier_key)}
            for key, price_key, profit_key, margin_key in (
                ("per_unit", "price_per_unit", "profit_per_unit", "margin_pct"),
                ("low", "price_low", "profit_low", "margin_low_pct"),
                ("high", "price_high", "profit_high", "margin_high_pct"),
            ):
                price = tier_prices.get(key)
                if price is not None:
                    margin = (price - total_per_unit) / price if price > 0 else 0
                    t[price_key] = price
                    t[profit_key] = round(price - total_per_unit, 4) if complete else None
                    t[margin_key] = round(margin * 100, 1) if complete else None
            entry["tiers"][tier_key] = t

        results.append(entry)
    return results


# ---------------------------------------------------------------------------
# Workbook output (same tabs as the Google Sheets version)
# ---------------------------------------------------------------------------
def write_workbook(path: str, ledger: list[list], unmapped: list[list],
                   prices: dict, costs: list[dict], config: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    tabs = tab_names(config)
    unit = config.get("products", {}).get("unit_name", "unit").title()
    pricing_tiers = config.get("pricing", {}).get("tiers", {})
    tier_keys = list(pricing_tiers.keys())
    overhead_names = [c["name"] for c in config.get("overhead", {}).get("cost_categories", [])]

    wb = Workbook()
    wb.remove(wb.active)

    def add(name, header, rows):
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for r in rows:
            ws.append(r)
        # Receipt text is data even when it begins with an Excel formula marker.
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)

    add(tabs["purchases"], PURCHASES_COLUMNS, ledger)

    add(tabs["latest_prices"],
        ["Item", "Unit Price", "Unit", "Qty Purchased", "Total Price", "Date", "Supplier"],
        [[n, i["unit_price"], i["unit"], i["quantity_purchased"], i["total_price"], i["date"], i["supplier"]]
         for n, i in sorted(prices.items())])

    def margin_text(value):
        return f"{value}%" if value is not None else ""

    def issues(result):
        return "; ".join(["Missing price: " + n for n in result["missing_prices"]] +
                         [e["ingredient"] + ": " + e["error"] for e in result["unit_errors"]] +
                         result["validation_errors"])

    header = ["Product", "Size", "Batch", f"Ingredient $/{unit}"]
    header += [f"{n} $/{unit}" for n in overhead_names] + [f"Total $/{unit}"]
    if tier_keys:
        label = pricing_tiers[tier_keys[0]].get("label", tier_keys[0])
        header += [f"{label} $/{unit}", f"{label} Margin %"]
    rows = []
    for r in costs:
        row = [r["product"], r["size"], r["batch_size"], r["ingredient_cost_per_unit"]]
        row += [r["overhead_breakdown"].get(n, 0) for n in overhead_names] + [r["total_cost_per_unit"]]
        if tier_keys:
            t = r["tiers"].get(tier_keys[0], {})
            row += [t.get("price_per_unit", ""), margin_text(t.get("margin_pct"))]
        row += [r["status"], issues(r)]
        rows.append(row)
    header += ["Status", "Issues"]
    add(tabs["recipes"], header, rows)

    header = ["Product", f"Total Cost/{unit}"]
    for k in tier_keys:
        label = pricing_tiers[k].get("label", k)
        header += [f"{label} Price", f"{label} Margin %", f"{label} Profit",
                   f"{label} High Price", f"{label} High Margin %", f"{label} High Profit"]
    header += ["Status", "Issues"]
    rows = []
    for r in costs:
        row = [r["product"], r["total_cost_per_unit"]]
        for k in tier_keys:
            t = r["tiers"].get(k, {})
            row += [t.get("price_per_unit", t.get("price_low", "")),
                    margin_text(t.get('margin_pct', t.get('margin_low_pct'))),
                    t.get("profit_per_unit", t.get("profit_low", "")),
                    t.get("price_high", ""), margin_text(t.get("margin_high_pct")),
                    t.get("profit_high", "")]
        row += [r["status"], issues(r)]
        rows.append(row)
    add(tabs["margins"], header, rows)

    add(tabs["unmapped"], PURCHASES_COLUMNS, unmapped)
    wb.save(path)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def run(config_path: str, receipt_paths: list[str], ledger_path: str | None,
        out_path: str | None, source: str = "chatgpt") -> dict:
    with open(config_path) as f:
        config = json.load(f)

    ledger = load_ledger(ledger_path, config)
    unmapped = [r for r in ledger if r[4] == "Uncategorized"]

    receipts, errors = [], []
    for path in receipt_paths:
        with open(path) as f:
            receipt = json.load(f)
        try:
            rows, summary = receipt_rows(receipt, ledger, config, source)
        except ValueError as e:
            errors.append({"receipt": path, "error": str(e)})
            continue
        ledger.extend(rows)
        unmapped.extend(r for r in rows if r[4] == "Uncategorized")
        receipts.append(summary)

    prices = latest_prices(ledger)
    costs = compute_costs(prices, config)
    if out_path:
        write_workbook(out_path, ledger, unmapped, prices, costs, config)

    return {
        "receipts_recorded": receipts,
        "errors": errors,
        "ledger_rows": len(ledger),
        "product_costs": [
            {k: c[k] for k in ("product", "unit_name", "ingredient_cost_per_unit",
                               "total_overhead_per_unit", "total_cost_per_unit",
                               "missing_prices", "unit_errors", "validation_errors", "status", "tiers")}
            for c in costs
        ],
        "output": out_path,
    }


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--config", required=True, help="business_config.json")
    p.add_argument("--receipt", action="append", default=[], help="receipt JSON (repeatable)")
    p.add_argument("--ledger", help="existing purchases ledger (.xlsx or .csv)")
    p.add_argument("--out", help="output workbook path (.xlsx)")
    p.add_argument("--source", default="chatgpt", help="value for the Source column")
    p.add_argument("--map", nargs="+", metavar="TEXT", help="only preview item mapping for these descriptions")
    args = p.parse_args(argv)

    if args.map:
        with open(args.config) as f:
            config = json.load(f)
        print(json.dumps([{"description": d, **map_item(d, config)} for d in args.map], indent=2))
        return 0

    result = run(args.config, args.receipt, args.ledger, args.out, args.source)
    print(json.dumps(result, indent=2, default=str))
    return 1 if result["errors"] and not result["receipts_recorded"] else 0


if __name__ == "__main__":
    sys.exit(main())
